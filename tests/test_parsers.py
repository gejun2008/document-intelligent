from pathlib import Path

import docx
import openpyxl
import pymupdf
from pptx import Presentation
from pptx.util import Inches

from docintel.config import RoutingConfig
from docintel.parsers.pdf import analyze_page
from docintel.models import Route
from docintel.parsers.pdf import PageStats
from docintel.router import plan_docx, plan_pdf, plan_pptx, plan_xlsx, route_pdf_page

from conftest import png_bytes


def test_pdf_routes(sample_pdf: Path):
    tasks = plan_pdf(sample_pdf, "d", RoutingConfig())
    assert [t.route for t in tasks] == [Route.PDF_TEXT, Route.AZURE_DI, Route.PDF_TEXT]
    assert "Quarterly report page 1" in tasks[0].native_md


def test_pdf_route_rules():
    cfg = RoutingConfig()
    S = PageStats
    assert route_pdf_page(S(500, 0.5, 0.0), cfg) is Route.AZURE_DI      # 文本层乱码
    assert route_pdf_page(S(500, 1.0, 0.8), cfg) is Route.VLM_GROUNDED  # 图表页
    assert route_pdf_page(S(10, 1.0, 0.0), cfg) is Route.AZURE_DI       # 几乎没字
    assert route_pdf_page(S(500, 1.0, 0.0, math_chars=80), cfg) is Route.VLM_GROUNDED  # 公式页
    assert route_pdf_page(S(500, 1.0, 1.0, invisible_ratio=0.95), cfg) is Route.AZURE_DI  # 扫描件叠 OCR 层
    assert route_pdf_page(S(500, 1.0, 0.0, math_chars=3), cfg) is Route.PDF_TEXT
    assert route_pdf_page(S(500, 1.0, 0.0, tables=1), cfg) is Route.VLM_GROUNDED       # 表格页
    assert route_pdf_page(S(500, 1.0, 0.0, tables=1), RoutingConfig(table_route="azure")) is Route.AZURE_DI
    assert route_pdf_page(S(500, 1.0, 0.0, tables=1), RoutingConfig(table_route="text")) is Route.PDF_TEXT


def _table_pdf(path: Path, ruled: bool) -> Path:
    """第 1 页正文，第 2 页一张 4 行 3 列的表（ruled=False 时不画框线）。"""
    doc = pymupdf.open()
    body = doc.new_page()
    body.insert_text((72, 72), "Annual summary of regional sales performance and outlook.", fontsize=11)
    page = doc.new_page()
    page.insert_text((72, 60), "Table 1. Revenue by region in thousands of units sold.", fontsize=11)
    rows = [["Region", "Q1", "Q2"], ["North", "120", "135"], ["South", "98", "104"], ["West", "77", "81"]]
    x0, y0, w, h = 72, 80, 120, 24
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            rect = pymupdf.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
            if ruled:
                page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 6, rect.y1 - 8), cell, fontsize=10)
    doc.save(path)
    return path


def test_pdf_table_pages_go_to_vision(tmp_path: Path):
    pdf = _table_pdf(tmp_path / "t.pdf", ruled=True)
    with pymupdf.open(pdf) as d:
        assert analyze_page(d[0]).tables == 0 and analyze_page(d[1]).tables == 1
    tasks = plan_pdf(pdf, "d", RoutingConfig())
    assert [t.route for t in tasks] == [Route.PDF_TEXT, Route.VLM_GROUNDED]
    assert "North" in tasks[1].native_text and tasks[1].native_md is None

    azure = plan_pdf(pdf, "d", RoutingConfig(table_route="azure"))
    assert azure[1].route is Route.AZURE_DI and "North" in azure[1].native_text  # 文本层留作质检参考

    kept = plan_pdf(pdf, "d", RoutingConfig(table_route="text"))
    assert kept[1].route is Route.PDF_TEXT and "North|120" in kept[1].native_md.replace(" ", "")


def test_pdf_borderless_table_caught_by_markdown(tmp_path: Path, monkeypatch):
    import docintel.router as router
    pdf = _table_pdf(tmp_path / "t.pdf", ruled=False)
    # 无框线表格框线检测找不到；pymupdf4llm 如果抽出了表格，就按它兜底改道
    monkeypatch.setattr(router, "pages_markdown", lambda doc, idx: {
        i: ("| Region | Q1 |\n|---|---|\n| North | 120 |" if i == 1 else "Annual summary") for i in idx})
    tasks = plan_pdf(pdf, "d", RoutingConfig())
    assert [t.route for t in tasks] == [Route.PDF_TEXT, Route.VLM_GROUNDED]
    assert tasks[0].native_md == "Annual summary"


def test_xlsx_merged_hidden_and_chunks(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["区域", "季度", "收入"])
    ws.append(["华东", "Q1", 100])
    ws.append([None, "Q2", 120.0])
    ws.merge_cells("A2:A3")
    for i in range(5):
        ws.append(["华南", f"Q{i}", i])
    hidden = wb.create_sheet("内部")
    hidden.append(["secret"])
    hidden.sheet_state = "hidden"
    path = tmp_path / "s.xlsx"
    wb.save(path)

    cfg = RoutingConfig(max_rows_per_table=4)
    tasks = plan_xlsx(path, "d", cfg)
    assert len(tasks) == 2                      # 7 行数据按 4 行分块；隐藏 sheet 跳过
    first = tasks[0].native_md
    assert "| 华东 | Q2 | 120 |" in first       # 合并单元格填充，120.0 → 120
    assert tasks[1].native_md.count("| 区域 | 季度 | 收入 |") == 1  # 续块重复表头
    assert all("secret" not in t.native_md for t in tasks)


def test_docx_structure_and_images(tmp_path: Path):
    d = docx.Document()
    d.add_heading("年度报告", level=1)
    d.add_paragraph("这是正文第一段，介绍了整体经营情况。")
    d.add_heading("财务数据", level=2)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "指标", "数值"
    table.cell(1, 0).text, table.cell(1, 1).text = "收入", "3400"
    img = tmp_path / "chart.png"
    img.write_bytes(png_bytes("chart", size=(800, 400)))
    d.add_picture(str(img), width=Inches(4))
    path = tmp_path / "r.docx"
    d.save(path)

    cfg = RoutingConfig(min_image_bytes=100)
    tasks = plan_docx(path, "d", cfg)
    body, image = tasks
    assert body.route is Route.NATIVE and image.route is Route.IMAGE
    assert "# 年度报告" in body.native_md and "## 财务数据" in body.native_md
    assert "| 收入 | 3400 |" in body.native_md
    assert image.placeholder in body.native_md


def test_pptx_slides(tmp_path: Path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "项目进展"
    s1.placeholders[1].text_frame.text = "第一阶段完成"
    s1.placeholders[1].text_frame.add_paragraph().text = "第二阶段进行中"
    s1.notes_slide.notes_text_frame.text = "讲到这里停一下"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "数据"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text, tbl.cell(1, 0).text, tbl.cell(1, 1).text = "a", "b", "1", "2"
    path = tmp_path / "p.pptx"
    prs.save(path)

    tasks = plan_pptx(path, "d", None, RoutingConfig())
    assert [t.route for t in tasks] == [Route.NATIVE, Route.NATIVE]
    assert "## 项目进展" in tasks[0].native_md and "- 第二阶段进行中" in tasks[0].native_md
    assert "> 备注：讲到这里停一下" in tasks[0].native_md
    assert "| 1 | 2 |" in tasks[1].native_md

    rendered = plan_pptx(path, "d", tmp_path / "fake.pdf", RoutingConfig())
    assert rendered[1].route is Route.VLM_GROUNDED and rendered[1].page_index == 1
