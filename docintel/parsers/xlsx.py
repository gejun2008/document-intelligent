"""xlsx 原生解析：直接读单元格值，不走 PDF/图片。

注意 data_only=True 读的是 Excel 上次保存时缓存的公式结果；如果文件是程序生成、
从没被 Excel 打开计算过，公式格会是 None。这种文件先用 LibreOffice 转一遍
xlsx→xlsx（会重新计算）再解析。

表格上方的合并标题行单独成段，带横向合并的首行当分组表头，和下一行压平成 "上级 下级"。
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

MAX_TITLE_ROWS = 3


@dataclass
class SheetChunk:
    sheet: str
    chunk_index: int
    markdown: str
    text: str


@dataclass
class SheetImage:
    sheet: str
    index: int
    data: bytes


_DECIMALS = re.compile(r"0\.(0+)")


def _decimals(fmt: str) -> int:
    m = _DECIMALS.search(fmt)
    return len(m.group(1)) if m else 0


def display(cell: Cell) -> str:
    """按单元格的数字格式输出人看到的样子：千分位、百分比、小数位。"""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ", timespec="seconds").removesuffix(" 00:00:00")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    fmt = (cell.number_format or "General").split(";")[0]
    if "%" in fmt:
        return f"{value * 100:.{_decimals(fmt)}f}%"
    if "#,##0" in fmt:
        return f"{value:,.{_decimals(fmt)}f}"
    if _DECIMALS.fullmatch(fmt):
        return f"{value:.{_decimals(fmt)}f}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _cell_md(text: str) -> str:
    return text.replace("|", "\\|").replace("\r", "").replace("\n", "<br>").strip()


def _layout(ws: Worksheet, fill_merged: bool) -> tuple[list[list[str]], list[bool]]:
    """返回裁掉空行空列后的网格，以及每行是否含横向合并。"""
    rows = [[display(c) for c in row] for row in ws.iter_rows()]
    hmerged = [False] * len(rows)
    for rng in ws.merged_cells.ranges:
        if rng.min_row > len(rows):
            continue
        if rng.max_col > rng.min_col:
            for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
                hmerged[r] = True
        if fill_merged:
            top = rows[rng.min_row - 1][rng.min_col - 1]
            for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
                for c in range(rng.min_col - 1, min(rng.max_col, len(rows[r]))):
                    rows[r][c] = top
    keep_rows = [i for i, r in enumerate(rows) if any(v.strip() for v in r)]
    if not keep_rows:
        return [], []
    width = max(len(rows[i]) for i in keep_rows)
    grid = [rows[i] + [""] * (width - len(rows[i])) for i in keep_rows]
    keep_cols = [c for c in range(width) if any(r[c].strip() for r in grid)]
    return [[r[c] for c in keep_cols] for r in grid], [hmerged[i] for i in keep_rows]


def _split_header(grid: list[list[str]], hmerged: list[bool]) -> tuple[list[str], list[str], list[list[str]]]:
    """拆出 (标题段落, 表头, 数据行)。"""
    titles: list[str] = []
    i = 0
    width = len(grid[0])
    while i < min(MAX_TITLE_ROWS, len(grid) - 1) and width > 1:
        distinct = {v for v in grid[i] if v.strip()}
        if len(distinct) != 1:
            break
        titles.append(distinct.pop())
        i += 1
    header = grid[i]
    if hmerged[i] and i + 1 < len(grid):
        sub = grid[i + 1]
        header = [a if not b or b == a else f"{a} {b}" for a, b in zip(header, sub)]
        i += 1
    return titles, header, grid[i + 1:]


def _table(header: list[str], body: list[list[str]]) -> str:
    lines = ["| " + " | ".join(_cell_md(h) for h in header) + " |",
             "|" + "---|" * len(header)]
    lines += ["| " + " | ".join(_cell_md(v) for v in row) + " |" for row in body]
    return "\n".join(lines)


def parse_xlsx(path: Path, *, include_hidden: bool, fill_merged: bool,
               max_rows: int, min_image_bytes: int) -> tuple[list[SheetChunk], list[SheetImage]]:
    wb = openpyxl.load_workbook(path, data_only=True)  # 合并单元格需要非 read_only 模式
    chunks: list[SheetChunk] = []
    images: list[SheetImage] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" and not include_hidden:
            continue
        grid, hmerged = _layout(ws, fill_merged)
        if grid:
            titles, header, body = _split_header(grid, hmerged)
            for i in range(0, max(len(body), 1), max_rows):
                part = body[i:i + max_rows]
                heading = f"## {ws.title}" + (f"（第 {i // max_rows + 1} 部分）" if len(body) > max_rows else "")
                md = "\n\n".join([heading, *(titles if i == 0 else []), _table(header, part)])
                text = "\n".join([*titles, *(" ".join(r) for r in [header, *part])])
                chunks.append(SheetChunk(ws.title, i // max_rows, md, text))
        # 内嵌图片走 VLM；openpyxl 没有公开 API，取不到就跳过
        for n, img in enumerate(getattr(ws, "_images", [])):
            try:
                data = img._data()
            except Exception:
                continue
            if len(data) >= min_image_bytes:
                images.append(SheetImage(ws.title, n, data))
    return chunks, images
