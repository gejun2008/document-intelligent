"""合成测试集：先生成内容和金标 md，再渲染成 docx / xlsx / pptx / 原生 PDF / 扫描 PDF。

金标零成本，而且同一份内容有原生和扫描两个版本，可以直接量出走视觉路径损失了多少。

金标约定（也是"理想输出"的定义）：
- 标题用 #，列表用 "- " 加两空格缩进，表格用 md 表格，第一行是表头
- 页眉、页脚、页码不进金标
- xlsx：每个 sheet 是 "## sheet 名" 加表格；纵向合并的单元格每行都填值；
  两级表头压平成 "上级 下级"；表格上方的标题行单独成段；隐藏 sheet 不进金标
- pptx：每页 "## 标题"，备注写成 "> 备注：..."
"""

from __future__ import annotations

import datetime as dt
import html
import io
import json
import random
from dataclasses import dataclass, field
from pathlib import Path

import docx
import openpyxl
import pymupdf
from openpyxl.styles import Font
from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.util import Inches, Pt

# ---- 内容模型 ---------------------------------------------------------------


@dataclass
class Heading:
    level: int
    text: str


@dataclass
class Para:
    text: str


@dataclass
class Bullets:
    items: list[tuple[int, str]]  # (层级, 文本)


@dataclass
class Table:
    header: list[str]
    rows: list[list[str]]


Block = Heading | Para | Bullets | Table


@dataclass
class Spec:
    name: str
    kind: str
    lang: str
    difficulty: str
    features: dict = field(default_factory=dict)


# ---- 文本生成 ---------------------------------------------------------------

_ORGS = ["华东区域", "研发中心", "销售一部", "供应链部门", "海外事业部", "客户服务中心", "华南分公司", "数据平台组"]
_PROJECTS = ["新一代结算系统", "智能客服升级", "仓储自动化改造", "渠道数字化项目", "质量追溯平台", "成本优化计划"]
_FACTORS = ["原材料价格波动", "汇率变化", "物流成本上升", "季节性需求回落", "新产品线投入"]
_PEOPLE = ["张伟", "李娜", "王强", "刘洋", "陈静", "杨帆", "赵磊", "周敏"]
_EN_ORGS = ["The north region", "The R&D team", "Our logistics unit", "The online channel", "The APAC division"]
_EN_PROJECTS = ["the billing platform", "the warehouse upgrade", "the partner portal", "the data lake migration"]
_COMPANIES = ["星河科技", "远航制造", "蓝海物流", "恒信咨询"]


class Gen:
    def __init__(self, seed: int, lang: str):
        self.rng = random.Random(seed)
        self.lang = lang  # zh | en | mixed

    def money(self) -> str:
        return f"{self.rng.randint(120, 98000):,}"

    def pct(self) -> str:
        return f"{self.rng.uniform(0.5, 60):.1f}"

    def date(self) -> str:
        d = dt.date(2025, 1, 1) + dt.timedelta(days=self.rng.randint(0, 600))
        return d.isoformat()

    def _zh(self) -> str:
        r = self.rng
        templates = [
            f"{r.choice(_ORGS)}本季度实现营业收入 {self.money()} 万元，同比增长 {self.pct()}%。",
            f"{r.choice(_ORGS)}的毛利率为 {self.pct()}%，较上年同期{r.choice(['上升', '下降'])} {self.pct()} 个百分点。",
            f"截至 {self.date()}，{r.choice(_ORGS)}累计完成订单 {r.randint(30, 9000)} 笔，客户满意度达到 {self.pct()}%。",
            f"受{r.choice(_FACTORS)}影响，{r.choice(_ORGS)}的运营成本上升至 {self.money()} 万元。",
            f"下一阶段，{r.choice(_ORGS)}将重点推进{r.choice(_PROJECTS)}，预计在 {self.date()} 前完成交付。",
            f"{r.choice(_PROJECTS)}由{r.choice(_PEOPLE)}负责，目前已完成 {r.randint(10, 95)}% 的里程碑。",
        ]
        return r.choice(templates)

    def _en(self) -> str:
        r = self.rng
        templates = [
            f"{r.choice(_EN_ORGS)} reported revenue of {self.money()} thousand, up {self.pct()}% year over year.",
            f"{r.choice(_EN_ORGS)} closed {r.randint(30, 9000)} orders by {self.date()}.",
            f"Work on {r.choice(_EN_PROJECTS)} is {r.randint(10, 95)}% complete and on track for {self.date()}.",
            f"Operating costs for {r.choice(_EN_PROJECTS)} rose to {self.money()} thousand.",
        ]
        return r.choice(templates)

    def sentence(self) -> str:
        if self.lang == "en" or (self.lang == "mixed" and self.rng.random() < 0.4):
            return self._en()
        return self._zh()

    def paragraph(self, n: tuple[int, int] = (2, 5)) -> str:
        sep = " " if self.lang == "en" else ""
        return sep.join(self.sentence() for _ in range(self.rng.randint(*n)))

    def title(self) -> str:
        if self.lang == "en":
            return self.rng.choice(["Quarterly Review", "Project Status", "Operations Update", "Budget Summary"])
        return self.rng.choice(["经营情况概述", "重点项目进展", "财务数据分析", "风险与应对", "下一步工作计划", "区域业绩对比"])

    def table(self, rows: int, wide: bool = False) -> Table:
        r = self.rng
        kind = r.choice(["finance", "products", "schedule"])
        if kind == "finance":
            years = ["2023", "2024", "2025"] + (["2026E", "2027E", "CAGR"] if wide else [])
            header = ["指标", *years]
            names = ["营业收入", "营业成本", "毛利", "销售费用", "管理费用", "研发费用", "净利润", "经营现金流",
                     "资本开支", "应收账款", "存货", "资产负债率"]
            body = [[names[i % len(names)] + ("" if i < len(names) else f"（{i // len(names) + 1}）"),
                     *[self.money() for _ in years[:-1]], f"{self.pct()}%"] for i in range(rows)]
        elif kind == "products":
            header = ["名称", "型号", "数量", "单价", "金额"] + (["仓库", "供应商", "备注"] if wide else [])
            body = []
            for i in range(rows):
                qty, price = r.randint(1, 500), r.randint(10, 9000)
                row = [f"{r.choice(['控制器', '传感器', '电源模块', '服务器', '交换机', '显示屏'])}-{i + 1}",
                       f"X{r.randint(100, 999)}", str(qty), f"{price:,}", f"{qty * price:,}"]
                if wide:
                    row += [r.choice(["上海", "苏州", "深圳"]), r.choice(_COMPANIES), r.choice(["", "加急", "含税"])]
                body.append(row)
        else:
            header = ["阶段", "负责人", "开始日期", "结束日期", "状态"] + (["预算（万元）", "风险等级"] if wide else [])
            body = []
            for i in range(rows):
                row = [f"阶段 {i + 1}", r.choice(_PEOPLE), self.date(), self.date(),
                       r.choice(["已完成", "进行中", "未开始", "延期"])]
                if wide:
                    row += [self.money(), r.choice(["低", "中", "高"])]
                body.append(row)
        return Table(header, body)

    def bullets(self, n: int, nested: bool) -> Bullets:
        items = []
        for _ in range(n):
            items.append((0, self.sentence()))
            if nested and self.rng.random() < 0.4:
                items.append((1, self.sentence()))
        return Bullets(items)

    def document(self, difficulty: str) -> list[Block]:
        r = self.rng
        sections = {"easy": 2, "medium": 3, "hard": 5}[difficulty]
        blocks: list[Block] = [Heading(1, f"{r.choice(_COMPANIES)}{r.choice(['年度报告', '季度经营分析', '项目周报'])}"),
                               Para(self.paragraph())]
        for s in range(sections):
            blocks.append(Heading(2, f"{s + 1}. {self.title()}"))
            blocks.append(Para(self.paragraph()))
            if difficulty != "easy" and r.random() < 0.6:
                blocks.append(self.bullets(r.randint(2, 5), nested=difficulty == "hard"))
            if r.random() < {"easy": 0.5, "medium": 0.7, "hard": 0.9}[difficulty]:
                rows = {"easy": (3, 5), "medium": (5, 12), "hard": (15, 45)}[difficulty]
                blocks.append(self.table(r.randint(*rows), wide=difficulty == "hard" and r.random() < 0.5))
            if difficulty == "hard":
                blocks.append(Heading(3, self.title()))
                blocks.append(Para(self.paragraph((3, 6))))
        return blocks


# ---- 渲染：金标 md ----------------------------------------------------------


def _md_cell(text: str) -> str:
    return text.replace("|", "\\|")


def md_table(t: Table) -> str:
    lines = ["| " + " | ".join(map(_md_cell, t.header)) + " |", "|" + "---|" * len(t.header)]
    lines += ["| " + " | ".join(map(_md_cell, row)) + " |" for row in t.rows]
    return "\n".join(lines)


def blocks_to_md(blocks: list[Block]) -> str:
    out = []
    for b in blocks:
        if isinstance(b, Heading):
            out.append("#" * b.level + " " + b.text)
        elif isinstance(b, Para):
            out.append(b.text)
        elif isinstance(b, Bullets):
            out.append("\n".join("  " * lvl + "- " + t for lvl, t in b.items))
        else:
            out.append(md_table(b))
    return "\n\n".join(out) + "\n"


# ---- 渲染：docx ---------------------------------------------------------------


def write_docx(blocks: list[Block], path: Path) -> None:
    d = docx.Document()
    for b in blocks:
        if isinstance(b, Heading):
            d.add_heading(b.text, level=b.level)
        elif isinstance(b, Para):
            d.add_paragraph(b.text)
        elif isinstance(b, Bullets):
            for lvl, t in b.items:
                d.add_paragraph(t, style="List Bullet" if lvl == 0 else "List Bullet 2")
        else:
            table = d.add_table(rows=len(b.rows) + 1, cols=len(b.header), style="Table Grid")
            for c, h in enumerate(b.header):
                cell = table.cell(0, c)
                cell.text = h
                cell.paragraphs[0].runs[0].bold = True
            for r, row in enumerate(b.rows, start=1):
                for c, v in enumerate(row):
                    table.cell(r, c).text = v
    d.save(path)


# ---- 渲染：PDF（原生 + 扫描）------------------------------------------------

_CJK = pymupdf.Font("cjk")
_CSS = """
@font-face {font-family: cjk; src: url(cjk.ttf);}
* {font-family: cjk; font-size: 10pt; line-height: 1.4;}
h1 {font-size: 18pt;} h2 {font-size: 14pt;} h3 {font-size: 12pt;}
table {border-collapse: collapse; margin: 6pt 0;}
td, th {border: 0.6pt solid #333; padding: 2pt 4pt; font-size: 9pt;}
th {background-color: #e8e8e8;}
"""


def blocks_to_html(blocks: list[Block]) -> str:
    e = html.escape
    out = []
    for b in blocks:
        if isinstance(b, Heading):
            out.append(f"<h{b.level}>{e(b.text)}</h{b.level}>")
        elif isinstance(b, Para):
            out.append(f"<p>{e(b.text)}</p>")
        elif isinstance(b, Bullets):
            parts, depth = ["<ul>"], 0
            for lvl, t in b.items:
                while depth < lvl:
                    parts.append("<ul>")
                    depth += 1
                while depth > lvl:
                    parts.append("</ul>")
                    depth -= 1
                parts.append(f"<li>{e(t)}</li>")
            parts.append("</ul>" * (depth + 1))
            out.append("".join(parts))
        else:
            head = "".join(f"<th>{e(h)}</th>" for h in b.header)
            body = "".join("<tr>" + "".join(f"<td>{e(v)}</td>" for v in row) + "</tr>" for row in b.rows)
            out.append(f"<table><tr>{head}</tr>{body}</table>")
    return "\n".join(out)


def write_pdf(blocks: list[Block], path: Path, header: str) -> int:
    archive = pymupdf.Archive()
    archive.add(_CJK.buffer, "cjk.ttf")
    story = pymupdf.Story(blocks_to_html(blocks), user_css=_CSS, archive=archive)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    page_rect = pymupdf.paper_rect("a4")
    content = page_rect + (56, 64, -56, -64)
    more = 1
    while more:
        dev = writer.begin_page(page_rect)
        more, _ = story.place(content)
        story.draw(dev)
        writer.end_page()
    writer.close()

    # 页眉页脚另外画：它们不进金标，用来测去页眉页脚
    doc = pymupdf.open("pdf", buf.getvalue())
    n = len(doc)
    for i, page in enumerate(doc):
        page.insert_font(fontname="cjk", fontbuffer=_CJK.buffer)
        page.insert_text((56, 40), header, fontname="cjk", fontsize=8, color=(0.4, 0.4, 0.4))
        page.insert_text((page_rect.width / 2 - 30, page_rect.height - 30), f"第 {i + 1} 页 共 {n} 页",
                         fontname="cjk", fontsize=8, color=(0.4, 0.4, 0.4))
    doc.save(path, garbage=3, deflate=True)
    doc.close()
    return n


def write_scanned_pdf(src: Path, dst: Path, rng: random.Random, dpi: int = 150) -> None:
    """模拟扫描件：灰度、轻微旋转、模糊、噪点、JPEG 压缩，最后只剩图片没有文本层。"""
    with pymupdf.open(src) as doc, pymupdf.open() as out:
        for page in doc:
            pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
            img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
            img = img.rotate(rng.uniform(-1.2, 1.2), resample=Image.BICUBIC, expand=False, fillcolor=255)
            img = img.filter(ImageFilter.GaussianBlur(rng.uniform(0.3, 0.9)))
            noise = Image.effect_noise(img.size, rng.uniform(8, 20))
            img = Image.blend(img, noise, rng.uniform(0.04, 0.1))
            jpg = io.BytesIO()
            img.save(jpg, format="JPEG", quality=rng.randint(45, 75))
            p = out.new_page(width=page.rect.width, height=page.rect.height)
            p.insert_image(p.rect, stream=jpg.getvalue())
        out.save(dst, garbage=3, deflate=True)


# ---- 渲染：xlsx ---------------------------------------------------------------


def write_xlsx(gen: Gen, difficulty: str, path: Path) -> tuple[str, dict]:
    r = gen.rng
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    gold: list[str] = []
    features = {"sheets": 0, "merged": 0, "hidden": 0, "title_rows": 0, "grouped_headers": 0}
    for s in range(r.randint(1, 3) if difficulty != "easy" else 1):
        name = f"{gen.title()[:6]}{s + 1}"
        ws = wb.create_sheet(name)
        features["sheets"] += 1
        rows = {"easy": (4, 8), "medium": (10, 40), "hard": (40, 200)}[difficulty]
        t = gen.table(r.randint(*rows))
        md_parts = [f"## {name}"]
        row_no = 1

        if difficulty == "hard" and r.random() < 0.6:  # 合并的标题行
            title = f"{r.choice(_COMPANIES)} {name} 明细表"
            ws.cell(row_no, 1, title).font = Font(bold=True, size=14)
            ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=len(t.header))
            md_parts.append(title)
            features["title_rows"] += 1
            row_no += 2

        header = list(t.header)
        if difficulty == "hard" and len(header) >= 5 and r.random() < 0.5:  # 两级表头
            ws.cell(row_no, 1, header[0])
            ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no + 1, end_column=1)
            ws.cell(row_no, 2, "基本信息")
            ws.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=3)
            ws.cell(row_no, 4, "数值")
            ws.merge_cells(start_row=row_no, start_column=4, end_row=row_no, end_column=len(header))
            for c, h in enumerate(header[1:], start=2):
                ws.cell(row_no + 1, c, h)
            header = [header[0]] + [f"基本信息 {h}" for h in header[1:3]] + [f"数值 {h}" for h in header[3:]]
            features["grouped_headers"] += 1
            features["merged"] += 3
            row_no += 2
        else:
            for c, h in enumerate(header, start=1):
                ws.cell(row_no, c, h).font = Font(bold=True)
            row_no += 1

        body = [list(row) for row in t.rows]
        merges: list[tuple[int, int]] = []
        if difficulty != "easy" and len(body) >= 6 and r.random() < 0.5:  # 纵向合并的分组列
            for start in range(0, len(body) - 2, r.randint(3, 5)):
                end = min(start + 2, len(body) - 1)
                for i in range(start, end + 1):
                    body[i][0] = body[start][0]
                merges.append((start, end))
        children = {i for s, e in merges for i in range(s + 1, e + 1)}
        for i, row in enumerate(body):
            for c, v in enumerate(row, start=1):
                if not (c == 1 and i in children):
                    _put(ws.cell(row_no + i, c), v)
        for s, e in merges:  # 先写值再合并，合并后的子格是只读的
            ws.merge_cells(start_row=row_no + s, start_column=1, end_row=row_no + e, end_column=1)
            features["merged"] += 1
        md_parts.append(md_table(Table(header, body)))
        gold.append("\n\n".join(md_parts))

    if difficulty != "easy" and r.random() < 0.5:
        hidden = wb.create_sheet("内部备注")
        hidden.append(["这一页是隐藏的，不应该出现在输出里"])
        hidden.sheet_state = "hidden"
        features["hidden"] += 1
    wb.save(path)
    return "\n\n".join(gold) + "\n", features


def _put(cell, v: str) -> None:
    """数字、百分比、日期写成真正的数值单元格加显示格式，金标是人在 Excel 里看到的样子。"""
    plain = v.replace(",", "")
    if plain.isdigit():
        cell.value = int(plain)
        if "," in v:
            cell.number_format = "#,##0"
    elif v.endswith("%") and plain[:-1].replace(".", "", 1).isdigit():
        cell.value = round(float(plain[:-1]) / 100, 6)
        cell.number_format = "0." + "0" * len(plain[:-1].split(".")[1]) + "%" if "." in plain else "0%"
    else:
        try:
            cell.value = dt.date.fromisoformat(v)
            cell.number_format = "yyyy-mm-dd"
        except ValueError:
            cell.value = v


# ---- 渲染：pptx ---------------------------------------------------------------


def write_pptx(gen: Gen, difficulty: str, path: Path) -> tuple[str, dict]:
    r = gen.rng
    prs = Presentation()
    gold = []
    n_slides = {"easy": 3, "medium": 6, "hard": 10}[difficulty]
    features = {"slides": n_slides, "tables": 0, "notes": 0}
    for _ in range(n_slides):
        title = gen.title()
        parts = [f"## {title}"]
        if r.random() < 0.35:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = title
            t = gen.table(r.randint(3, 7 if difficulty == "hard" else 5))
            shape = slide.shapes.add_table(len(t.rows) + 1, len(t.header), Inches(0.5), Inches(1.6),
                                           Inches(9), Inches(0.3) * (len(t.rows) + 1))
            for c, h in enumerate(t.header):
                shape.table.cell(0, c).text = h
            for i, row in enumerate(t.rows, start=1):
                for c, v in enumerate(row):
                    shape.table.cell(i, c).text = v
            for row in shape.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(10)
            parts.append(md_table(t))
            features["tables"] += 1
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            b = gen.bullets(r.randint(2, 4), nested=difficulty != "easy")
            tf = slide.placeholders[1].text_frame
            for i, (lvl, text) in enumerate(b.items):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text, p.level = text, lvl
                p.font.size = Pt(14)
            parts.append("\n".join("  " * lvl + "- " + t for lvl, t in b.items))
        if difficulty != "easy" and r.random() < 0.4:
            note = gen.sentence()
            slide.notes_slide.notes_text_frame.text = note
            parts.append(f"> 备注：{note}")
            features["notes"] += 1
        gold.append("\n\n".join(parts))
    prs.save(path)
    return "\n\n".join(gold) + "\n", features


# ---- 入口 -------------------------------------------------------------------

KINDS = ("docx", "xlsx", "pptx", "pdf_native", "pdf_scan")
DIFFICULTIES = ("easy", "medium", "hard")
LANGS = ("zh", "zh", "mixed", "en")  # 中文为主


def _count(blocks: list[Block]) -> dict:
    tables = [b for b in blocks if isinstance(b, Table)]
    return {"tables": len(tables), "table_rows": sum(len(t.rows) for t in tables),
            "headings": sum(isinstance(b, Heading) for b in blocks)}


def generate(out: Path, per_kind: int = 9, seed: int = 0) -> list[Spec]:
    docs, gold = out / "docs", out / "gold"
    docs.mkdir(parents=True, exist_ok=True)
    gold.mkdir(parents=True, exist_ok=True)
    specs: list[Spec] = []

    for i in range(per_kind):
        difficulty = DIFFICULTIES[i % len(DIFFICULTIES)]
        for kind in ("docx", "xlsx", "pptx", "pdf"):
            item_seed = seed * 100_000 + i * 10 + KINDS.index(kind if kind != "pdf" else "pdf_native")
            lang = LANGS[(i + item_seed) % len(LANGS)]
            gen = Gen(item_seed, lang)
            if kind == "docx":
                blocks = gen.document(difficulty)
                name = f"docx_{i:03d}"
                write_docx(blocks, docs / f"{name}.docx")
                (gold / f"{name}.md").write_text(blocks_to_md(blocks), "utf-8")
                specs.append(Spec(name, "docx", lang, difficulty, _count(blocks)))
            elif kind == "xlsx":
                name = f"xlsx_{i:03d}"
                md, feats = write_xlsx(gen, difficulty, docs / f"{name}.xlsx")
                (gold / f"{name}.md").write_text(md, "utf-8")
                specs.append(Spec(name, "xlsx", lang, difficulty, feats))
            elif kind == "pptx":
                name = f"pptx_{i:03d}"
                md, feats = write_pptx(gen, difficulty, docs / f"{name}.pptx")
                (gold / f"{name}.md").write_text(md, "utf-8")
                specs.append(Spec(name, "pptx", lang, difficulty, feats))
            else:
                blocks = gen.document(difficulty)
                md = blocks_to_md(blocks)
                native, scan = f"pdf_native_{i:03d}", f"pdf_scan_{i:03d}"
                pages = write_pdf(blocks, docs / f"{native}.pdf", header=f"{gen.rng.choice(_COMPANIES)} · 内部资料")
                write_scanned_pdf(docs / f"{native}.pdf", docs / f"{scan}.pdf", gen.rng)
                for name, k in ((native, "pdf_native"), (scan, "pdf_scan")):
                    (gold / f"{name}.md").write_text(md, "utf-8")
                    specs.append(Spec(name, k, lang, difficulty, {**_count(blocks), "pages": pages}))

    with open(out / "manifest.jsonl", "w", encoding="utf-8") as f:
        for s in specs:
            f.write(json.dumps(s.__dict__, ensure_ascii=False) + "\n")
    return specs
