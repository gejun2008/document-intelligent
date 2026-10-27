"""合成测试集上发现的问题的回归测试。"""

import asyncio
from pathlib import Path

import openpyxl

from docintel.assemble import strip_headers_footers
from docintel.config import RoutingConfig
from docintel.pipeline import Pipeline
from docintel.quality.benchmark import run_benchmark
from docintel.quality.checks import repetition_ratio, split_row, table_consistency
from docintel.router import plan_xlsx
from docintel.synth import generate


def test_xlsx_title_row_grouped_header_and_number_formats(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "2025 年销售明细"
    ws.merge_cells("A1:D1")
    ws["A3"], ws["B3"], ws["D3"] = "区域", "上半年", "合计"
    ws.merge_cells("A3:A4")
    ws.merge_cells("B3:C3")
    ws.merge_cells("D3:D4")
    ws["B4"], ws["C4"] = "Q1", "Q2"
    ws.append(["华东", 12345, 0.125, 1234.5])
    ws["B5"].number_format = "#,##0"
    ws["C5"].number_format = "0.0%"
    ws["D5"].number_format = "#,##0.00"
    path = tmp_path / "t.xlsx"
    wb.save(path)

    md = plan_xlsx(path, "d", RoutingConfig())[0].native_md
    assert "\n\n2025 年销售明细\n\n" in md                    # 标题行单独成段，不进表头
    assert "| 区域 | 上半年 Q1 | 上半年 Q2 | 合计 |" in md    # 两级表头压平
    assert "| 华东 | 12,345 | 12.5% | 1,234.50 |" in md      # 按显示格式输出


def test_split_row_keeps_trailing_empty_cell():
    assert split_row("|a|b||") == ["a", "b", ""]
    assert table_consistency("|a|b|c|\n|---|---|---|\n|1|2||\n|3|4|5|")[0] == 1.0


def test_repetition_ignores_legit_table_repeats_but_catches_row_loops():
    filled = "| 华东 | 华东 | 华东 | 华东 | 华东 | 华东 |\n|---|---|---|---|---|---|\n" + "| 低 | 低 | 低 | 低 | 低 | 低 |\n" * 3
    assert repetition_ratio(filled) == 0.0
    loop = "| a | b |\n|---|---|\n" + "| 同一行 | 123 |\n" * 6
    assert repetition_ratio(loop) == 1.0


def test_page_numbers_stripped_even_for_short_docs():
    assert strip_headers_footers(["正文第一段\n\n第 1 页 共 1 页"]) == ["正文第一段"]
    assert strip_headers_footers(["Body text\n\nPage 2 of 9"]) == ["Body text"]


def test_synthetic_native_paths_match_gold(tmp_path: Path, cfg):
    specs = generate(tmp_path / "ts", per_kind=3, seed=7)
    files = [tmp_path / "ts" / "docs" / f"{s.name}.{s.kind}" for s in specs if s.kind in ("docx", "xlsx", "pptx")]
    files += [tmp_path / "ts" / "docs" / f"{s.name}.pdf" for s in specs if s.kind == "pdf_native"]
    cfg.routing.table_route = "text"  # 这里只测文本层本身；表格页默认改走 VLM，本地没有模型

    async def go():
        async with Pipeline(cfg) as pipe:
            return await pipe.run(files)

    docs = asyncio.run(go())
    assert all(d.error is None for d in docs)
    _, summary = run_benchmark(cfg.out_dir, tmp_path / "ts" / "gold", {f.stem for f in files})
    for kind in ("docx", "xlsx", "pptx"):
        assert summary[kind]["text_ned"] == 0 and summary[kind]["table_cell_f1"] == 1, (kind, summary[kind])
    assert summary["pdf_native"]["number_f1"] > 0.99 and summary["pdf_native"]["heading_f1"] > 0.9
