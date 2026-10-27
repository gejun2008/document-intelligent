"""pptx 原生解析：抽文本框、表格、备注。

结果有两个用途：没有渲染条件时直接当输出；有渲染时作为 grounded VLM 的文字提示，
让 VLM 负责版式和阅读顺序、文字以原生为准。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass
class SlideContent:
    index: int          # 在所有幻灯片里的序号（含隐藏）
    visible_index: int  # 在可见幻灯片里的序号，对应 LibreOffice 导出 PDF 的页号
    markdown: str
    text: str


def _cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", "<br>").strip()


def _shape_lines(shape, title_id: int | None) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        lines: list[str] = []
        for sub in sorted(shape.shapes, key=lambda s: (s.top or 0, s.left or 0)):
            lines += _shape_lines(sub, title_id)
        return lines
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[_cell(c.text) for c in r.cells] for r in shape.table.rows]
        if not rows:
            return []
        out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
        out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return ["", *out, ""]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        if shape.shape_id == title_id:
            return [f"## {shape.text_frame.text.strip()}", ""]
        lines = []
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                lines.append("  " * p.level + f"- {t}" if p.level or len(shape.text_frame.paragraphs) > 1 else t)
        return [*lines, ""] if lines else []
    return []


def parse_pptx(path: Path) -> list[SlideContent]:
    prs = Presentation(str(path))
    slides: list[SlideContent] = []
    visible = 0
    for i, slide in enumerate(prs.slides):
        if slide._element.get("show") == "0":  # 隐藏幻灯片 LibreOffice 默认不导出
            continue
        title = slide.shapes.title
        title_id = title.shape_id if title is not None else None
        lines: list[str] = []
        for shape in sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0)):
            lines += _shape_lines(shape, title_id)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines += ["", f"> 备注：{notes}"]
        md = "\n".join(lines).strip()
        text = "\n".join(l.lstrip("#-| >") for l in lines)
        slides.append(SlideContent(i, visible, md, text))
        visible += 1
    return slides
